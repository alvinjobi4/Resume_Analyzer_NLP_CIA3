import io
from typing import Union, Tuple, Dict, Any

def extract_text_from_docx(file_input: Union[str, bytes, io.BytesIO]) -> Tuple[str, Dict[str, Any]]:
    """
    Extract text from a DOCX file using python-docx.
    Extracts text from paragraphs and table cells to capture all resume details.
    Handles file paths, bytes, or Streamlit UploadedFile objects.
    """
    try:
        import docx
    except ImportError:
        raise ImportError("python-docx is required for DOCX parsing. Please install python-docx.")

    text_parts = []
    metadata = {
        "format": "DOCX",
        "paragraph_count": 0,
        "table_count": 0,
        "char_count": 0
    }

    try:
        if isinstance(file_input, str):
            doc = docx.Document(file_input)
        elif isinstance(file_input, bytes):
            doc = docx.Document(io.BytesIO(file_input))
        elif hasattr(file_input, "read"):
            stream_bytes = file_input.read()
            if hasattr(file_input, "seek"):
                file_input.seek(0)
            doc = docx.Document(io.BytesIO(stream_bytes))
        else:
            raise ValueError(f"Unsupported file input type: {type(file_input)}")

        # Extract paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                text_parts.append(text)
                metadata["paragraph_count"] += 1

        # Extract table contents (frequently used in resume layouts)
        for table in doc.tables:
            metadata["table_count"] += 1
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                # Deduplicate identical adjacent cells (due to merged cells in Word)
                deduped_cells = []
                for cell_text in row_cells:
                    if not deduped_cells or cell_text != deduped_cells[-1]:
                        deduped_cells.append(cell_text)
                if deduped_cells:
                    text_parts.append(" | ".join(deduped_cells))

        full_text = "\n".join(text_parts).strip()
        metadata["char_count"] = len(full_text)
        return full_text, metadata

    except Exception as e:
        raise RuntimeError(f"Error parsing DOCX file: {str(e)}") from e
