from pathlib import Path
import fitz  # PyMuPDF
import docx

def create_sample_pdf(text: str, output_path: Path):
    """Generates a clean test PDF resume using PyMuPDF."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)  # A4 size
    rect = fitz.Rect(50, 50, 545, 792)
    page.insert_textbox(rect, text, fontsize=10, fontname="helv", align=0)
    doc.save(str(output_path))
    doc.close()
    print(f"Created PDF: {output_path}")

def create_sample_docx(text: str, output_path: Path):
    """Generates a test DOCX resume using python-docx."""
    doc = docx.Document()
    for line in text.split("\n"):
        line_clean = line.strip()
        if not line_clean:
            continue
        if line_clean in ["SUMMARY", "EDUCATION", "TECHNICAL SKILLS", "WORK EXPERIENCE", "PROJECTS", "CERTIFICATIONS"]:
            doc.add_heading(line_clean, level=2)
        elif line_clean.startswith("John Doe"):
            doc.add_heading(line_clean, level=1)
        else:
            doc.add_paragraph(line_clean)
    doc.save(str(output_path))
    print(f"Created DOCX: {output_path}")

if __name__ == "__main__":
    base = Path(__file__).resolve().parent
    txt_path = base / "sample_resume.txt"
    if txt_path.exists():
        with open(txt_path, "r", encoding="utf-8") as f:
            content = f.read()
        create_sample_pdf(content, base / "sample_resume.pdf")
        create_sample_docx(content, base / "sample_resume.docx")
        print("Sample test files generated successfully!")
